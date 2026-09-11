#!/usr/bin/env python3
"""Create compact timing and resource reports from an RDF matrix summary."""
from __future__ import annotations

import argparse
import csv
import json
import os
import statistics
import sys
import tempfile
from pathlib import Path
from typing import Any, Mapping


def _number(value: Any, field: str) -> int:
    if not isinstance(value, int) or isinstance(value, bool) or value < 0:
        raise ValueError(f"{field} must be a non-negative integer")
    return value


def _validate_timing(value: Mapping[str, Any], schema: str, field: str) -> None:
    if not isinstance(value, Mapping) or value.get("schema") != schema:
        raise ValueError(f"{field} has an invalid schema")
    stages = value.get("stages_ns")
    total = _number(value.get("total_wall_ns"), f"{field}.total_wall_ns")
    if not isinstance(stages, Mapping):
        raise ValueError(f"{field}.stages_ns must be an object")
    stage_sum = sum(_number(item, f"{field}.stages_ns") for item in stages.values())
    if stage_sum != total or value.get("stages_sum_ns", stage_sum) != stage_sum:
        raise ValueError(f"{field} does not reconcile")
    if value.get("reconciled") is not True:
        raise ValueError(f"{field} is not marked reconciled")


def _milliseconds(value: int | None) -> float | None:
    return None if value is None else round(value / 1_000_000, 3)


def _phase(workload: Mapping[str, Any], name: str) -> Mapping[str, Any]:
    phases = workload.get("phases", {})
    value = phases.get(name, {})
    return value if isinstance(value, Mapping) else {}


def _phase_memory(value: Any, field: str) -> Mapping[str, Any] | None:
    if value is None:
        return None
    if not isinstance(value, Mapping) or value.get("schema") != "rdf-phase-memory-metrics-v1":
        raise ValueError(f"{field} has an invalid phase memory schema")
    if value.get("unit") != "bytes":
        raise ValueError(f"{field} must use bytes")
    scope = value.get("scope")
    if not isinstance(scope, str) or not scope:
        raise ValueError(f"{field}.scope must be a non-empty string")
    for name in ("sample_count", "sample_errors"):
        _number(value.get(name), f"{field}.{name}")
    peak = value.get("peak_rss_bytes")
    if peak is not None:
        _number(peak, f"{field}.peak_rss_bytes")
    interval = value.get("sampling_interval_ms")
    if not isinstance(interval, (int, float)) or isinstance(interval, bool) or interval <= 0:
        raise ValueError(f"{field}.sampling_interval_ms must be positive")
    phases = value.get("phases")
    if not isinstance(phases, Mapping):
        raise ValueError(f"{field}.phases must be an object")
    for phase, metrics in phases.items():
        if not isinstance(phase, str) or not isinstance(metrics, Mapping):
            raise ValueError(f"{field}.phases is invalid")
        _number(metrics.get("sample_count"), f"{field}.phases.{phase}.sample_count")
        for name in ("first_rss_bytes", "last_rss_bytes", "peak_rss_bytes"):
            item = metrics.get(name)
            if item is not None:
                _number(item, f"{field}.phases.{phase}.{name}")
    return value


def _mebibytes(value: int | None) -> float | None:
    return None if value is None else round(value / (1024 * 1024), 3)


def _memory_phase_peak(memory: Mapping[str, Any] | None, phase: str) -> float | None:
    if memory is None:
        return None
    value = memory["phases"].get(phase)
    if not isinstance(value, Mapping):
        return None
    return _mebibytes(value.get("peak_rss_bytes"))


def build_report(summary: Mapping[str, Any]) -> dict[str, Any]:
    experiments = summary.get("experiments")
    if not isinstance(experiments, list):
        raise ValueError("summary experiments must be an array")
    rows = []
    for index, experiment in enumerate(experiments):
        if not isinstance(experiment, Mapping):
            raise ValueError(f"experiment {index} must be an object")
        timing = experiment.get("system_timing")
        _validate_timing(timing, "rdf-system-timing-v1", f"experiment {index}.system_timing")
        workload = experiment.get("workload_timing")
        if not isinstance(workload, Mapping) or workload.get("schema") != "rdf-workload-timing-v1":
            raise ValueError(f"experiment {index} has invalid workload timing")
        warmup = _phase(workload, "warmup")
        measured = _phase(workload, "measured")
        measured_count = _number(measured.get("attempt_count", 0), "measured.attempt_count")
        measured_total = _number(measured.get("attempt_total_ns", 0), "measured.attempt_total_ns")
        successful_count = _number(
            measured.get("successful_attempt_count", measured_count if experiment.get("failure_count", 0) == 0 else 0),
            "measured.successful_attempt_count",
        )
        successful_total = _number(
            measured.get("successful_attempt_total_ns", measured_total if experiment.get("failure_count", 0) == 0 else 0),
            "measured.successful_attempt_total_ns",
        )
        metrics = experiment.get("resource_metrics") or {}
        outcomes = experiment.get("outcome_counts") or {}
        build = experiment.get("build_metrics") or {}
        build_memory = build.get("memory") or {}
        representation_size = experiment.get("representation_size") or {}
        orchestration = experiment.get("runtime_orchestration") or {}
        memory = _phase_memory(
            experiment.get("phase_memory_metrics"),
            f"experiment {index}.phase_memory_metrics",
        )
        mode = experiment.get("execution_mode") or {}
        load_temperature = experiment.get("load_temperature_metrics") or {}
        stages = timing["stages_ns"]
        rows.append({
            "system": experiment.get("system"),
            "representation": experiment.get("representation"),
            "status": experiment.get("status"),
            "record_count": experiment.get("record_count"),
            "failure_count": experiment.get("failure_count"),
            "completed_count": outcomes.get("completed", experiment.get("success_count")),
            "timeout_count": outcomes.get("timeout"),
            "skipped_count": outcomes.get("skipped", experiment.get("skipped_count")),
            "engine_error_count": outcomes.get("engine-error"),
            "semantic_mismatch_count": outcomes.get("semantic-mismatch"),
            "confirmed_oom_count": outcomes.get("confirmed-oom"),
            "storage": mode.get("storage"),
            "process_temperature": mode.get("process_temperature"),
            "lifecycle": mode.get("lifecycle"),
            "warmup_total_ms": _milliseconds(warmup.get("attempt_total_ns", 0)),
            "measured_total_ms": _milliseconds(measured_total),
            "measured_attempt_count": measured_count,
            "successful_attempt_count": successful_count,
            "successful_query_total_ms": _milliseconds(successful_total),
            "successful_query_mean_ms": (
                None if successful_count == 0
                else round(successful_total / successful_count / 1_000_000, 3)
            ),
            "measured_total_ms": _milliseconds(measured_total),
            "measured_mean_ms": (
                None if measured_count == 0
                else round(measured_total / measured_count / 1_000_000, 3)
            ),
            "system_total_ms": _milliseconds(timing["total_wall_ns"]),
            "open_or_load_ms": _milliseconds(stages.get("artifact_open_or_load", 0)),
            "load_temperature_status": load_temperature.get("status"),
            "process_cold_load_or_parse_ms": _milliseconds(
                load_temperature.get("process_cold_load_or_parse_ns")
            ),
            "process_warm_load_or_parse_ms": _milliseconds(
                load_temperature.get("process_warm_load_or_parse_ns")
            ),
            "restart_load_or_parse_ms": _milliseconds(
                load_temperature.get("restart_load_or_parse_ns")
            ),
            "load_temperature_semantics": load_temperature.get("semantics"),
            "build_status": build.get("status"),
            "build_time_ms": _milliseconds(build.get("elapsed_ns")),
            "build_peak_rss_mib": _mebibytes(build_memory.get("peak_rss_bytes")),
            "representation_logical_bytes": representation_size.get("logical_bytes"),
            "representation_allocated_bytes": representation_size.get("allocated_bytes"),
            "representation_size_boundary": representation_size.get("boundary"),
            "matrix_run_id": orchestration.get("matrix_run_id"),
            "probe_policy_id": orchestration.get("probe_policy_id"),
            "probe_policy_sha256": orchestration.get("probe_policy_sha256"),
            "quarantine_snapshot_sha256": orchestration.get("quarantine_snapshot_sha256"),
            "result_file": experiment.get("result_file"),
            "startup_ms": _milliseconds(stages.get("engine_startup", 0)),
            "shutdown_ms": _milliseconds(stages.get("engine_shutdown", 0)),
            "validation_ms": _milliseconds(stages.get("validation", 0)),
            "unclassified_ms": _milliseconds(stages.get("unclassified", 0)),
            "max_rss_mib": (
                None if metrics.get("max_rss_kib") is None
                else round(_number(metrics["max_rss_kib"], "max_rss_kib") / 1024, 3)
            ),
            "memory_scope": None if memory is None else memory["scope"],
            "memory_sampling_interval_ms": (
                None if memory is None else memory["sampling_interval_ms"]
            ),
            "memory_sample_count": None if memory is None else memory["sample_count"],
            "memory_sample_errors": None if memory is None else memory["sample_errors"],
            "peak_rss_mib": (
                None if memory is None else _mebibytes(memory["peak_rss_bytes"])
            ),
            "open_or_load_peak_rss_mib": _memory_phase_peak(
                memory, "artifact_open_or_load"
            ),
            "warmup_peak_rss_mib": _memory_phase_peak(memory, "warmup"),
            "measured_peak_rss_mib": _memory_phase_peak(memory, "measured"),
            "shutdown_peak_rss_mib": _memory_phase_peak(
                memory, "engine_shutdown"
            ),
            "user_cpu_ms": _milliseconds(metrics.get("user_cpu_ns")),
            "system_cpu_ms": _milliseconds(metrics.get("system_cpu_ns")),
            "major_page_faults": metrics.get("major_page_faults"),
            "minor_page_faults": metrics.get("minor_page_faults"),
            "input_blocks": metrics.get("input_blocks"),
            "output_blocks": metrics.get("output_blocks"),
        })
    measured_values = [row["measured_total_ms"] for row in rows]
    return {
        "schema": "rdf-experiment-report-v1",
        "status": summary.get("status"),
        "matrix_timing": summary.get("matrix_timing"),
        "experiment_count": len(rows),
        "measured_total_ms": round(sum(measured_values), 3),
        "measured_median_ms": (
            None if not measured_values
            else round(statistics.median(measured_values), 3)
        ),
        "experiments": rows,
    }



def _atomic_write(path: Path, writer) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(prefix=f'.{path.name}.', suffix='.tmp', dir=path.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        writer(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_json(path: Path, report: Mapping[str, Any]) -> None:
    _atomic_write(path, lambda target: target.write_text(
        json.dumps(report, indent=2, sort_keys=True, allow_nan=False) + "\n", encoding="utf-8"
    ))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    def emit(target: Path) -> None:
        if not rows:
            target.write_text("", encoding="utf-8")
            return
        with target.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader(); writer.writerows(rows)
    _atomic_write(path, emit)


def write_markdown(path: Path, report: Mapping[str, Any]) -> None:
    rows = report["experiments"]
    columns = ["system", "representation", "status", "successful_attempt_count",
               "timeout_count", "skipped_count", "engine_error_count",
               "semantic_mismatch_count", "confirmed_oom_count",
               "successful_query_total_ms", "measured_peak_rss_mib",
               "process_cold_load_or_parse_ms", "process_warm_load_or_parse_ms",
               "build_time_ms", "build_peak_rss_mib", "representation_logical_bytes"]
    def cell(value: Any) -> str:
        if value is None: return "not measured"
        return str(value).replace("|", "\\|").replace("\n", " ")
    lines = ["# RDF experiment report", "", f"Status: **{cell(report.get('status'))}**", "",
             "| " + " | ".join(columns) + " |", "| " + " | ".join("---" for _ in columns) + " |"]
    lines.extend("| " + " | ".join(cell(row.get(name)) for name in columns) + " |" for row in rows)
    lines.extend(["", "Timeouts, errors, semantic mismatches, and confirmed OOM outcomes are not converted to successful query times.", ""])
    _atomic_write(path, lambda target: target.write_text("\n".join(lines), encoding="utf-8"))


def write_xlsx(path: Path, report: Mapping[str, Any]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError("XLSX output requires openpyxl; install execution-framework/requirements.txt") from error
    rows = report["experiments"]
    def emit(target: Path) -> None:
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Experiment summary"
        headers = list(rows[0]) if rows else ["status"]
        sheet.append(headers)
        for row in rows: sheet.append([row.get(name) for name in headers])
        fill=PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font=Font(bold=True,color="FFFFFF"); cell.fill=fill; cell.alignment=Alignment(wrap_text=True)
        sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
        for index,name in enumerate(headers,1):
            width=max(len(str(name)),*(len(str(row.get(name,''))) for row in rows)) if rows else len(name)
            sheet.column_dimensions[get_column_letter(index)].width=min(max(width+2,12),48)
        meta=workbook.create_sheet("Metadata")
        meta.append(["field","value"]); meta.append(["schema",report.get("schema")]); meta.append(["status",report.get("status")]); meta.append(["experiment_count",report.get("experiment_count")])
        for cell in meta[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=fill
        meta.column_dimensions["A"].width=24; meta.column_dimensions["B"].width=48
        workbook.save(target)
    _atomic_write(path, emit)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("summary", type=Path)
    parser.add_argument("--json", dest="json_path", type=Path)
    parser.add_argument("--csv", dest="csv_path", type=Path)
    parser.add_argument("--markdown", dest="markdown_path", type=Path)
    parser.add_argument("--xlsx", dest="xlsx_path", type=Path)
    args = parser.parse_args(argv)
    report = build_report(json.loads(args.summary.read_text(encoding="utf-8")))
    if args.json_path: write_json(args.json_path, report)
    if args.csv_path: write_csv(args.csv_path, report["experiments"])
    if args.markdown_path: write_markdown(args.markdown_path, report)
    if args.xlsx_path: write_xlsx(args.xlsx_path, report)
    if not any((args.json_path,args.csv_path,args.markdown_path,args.xlsx_path)):
        json.dump(report, sys.stdout, indent=2); sys.stdout.write("\n")
    return 0

if __name__ == "__main__": raise SystemExit(main())
