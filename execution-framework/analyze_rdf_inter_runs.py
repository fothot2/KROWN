#!/usr/bin/env python3
"""Create JSON, CSV, Markdown, and XLSX RDF inter-run statistics reports."""
from __future__ import annotations
import argparse,csv,json,os,tempfile
from pathlib import Path
from bench_executor.rdf_inter_run_statistics import build_inter_run_statistics,flatten_groups

def atomic(path,writer):
 path.parent.mkdir(parents=True,exist_ok=True); fd,name=tempfile.mkstemp(prefix='.'+path.name+'.',suffix='.tmp',dir=path.parent); os.close(fd); tmp=Path(name)
 try: writer(tmp); os.replace(tmp,path)
 finally: tmp.unlink(missing_ok=True)

def load(paths):
 records=[]; sources=[]
 for path in paths:
  with path.open(encoding='utf-8') as stream:
   for number,line in enumerate(stream,1):
    if line.strip(): records.append(json.loads(line)); sources.append(f'{path}:{number}')
 return records,sources

def write_csv(path,rows):
 def emit(target):
  with target.open('w',encoding='utf-8',newline='') as out:
   fields=list(rows[0]) if rows else ['record_count']; writer=csv.DictWriter(out,fieldnames=fields); writer.writeheader(); writer.writerows(rows)
 atomic(path,emit)

def write_markdown(path,rows):
 columns=['system','representation','query_id','phase','measurement_boundary','elapsed_ns_observation_count','elapsed_ns_mean','elapsed_ns_median','outcome_timeout','outcome_semantic-mismatch','outcome_confirmed-oom']
 def cell(v): return 'not measured' if v is None else str(v).replace('|','\\|')
 lines=['# RDF inter-run statistics','','| '+' | '.join(columns)+' |','| '+' | '.join('---' for _ in columns)+' |']
 lines += ['| '+' | '.join(cell(row.get(c)) for c in columns)+' |' for row in rows]
 lines += ['','Only completed outcomes contribute to elapsed-time statistics.','']
 atomic(path,lambda target:target.write_text('\n'.join(lines),encoding='utf-8'))

def write_xlsx(path,report,rows):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 from openpyxl.utils import get_column_letter
 def excel_value(value):
  if isinstance(value,(dict,list,tuple)):
   return json.dumps(value,sort_keys=True,separators=(',',':'))
  return value
 def add(ws,data):
  headers=list(data[0]) if data else ['record_count']; ws.append(headers)
  for row in data: ws.append([excel_value(row.get(h)) for h in headers])
  fill=PatternFill('solid',fgColor='1F4E78')
  for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=fill
  ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
  for i,h in enumerate(headers,1): ws.column_dimensions[get_column_letter(i)].width=min(max(len(h)+2,12),42)
 def emit(target):
  wb=Workbook(); ws=wb.active; ws.title='Inter-run statistics'; add(ws,rows)
  outcomes=[]
  for g in report['groups']: outcomes.append({**{k:g[k] for k in ('system','representation','query_id','phase')},**g['outcome_counts']})
  add(wb.create_sheet('Outcome counts'),outcomes); add(wb.create_sheet('Observation provenance'),report['observations']); wb.save(target)
 atomic(path,emit)

def main(argv=None):
 p=argparse.ArgumentParser(); p.add_argument('records',type=Path,nargs='+'); p.add_argument('--json',dest='json_path',type=Path); p.add_argument('--csv',dest='csv_path',type=Path); p.add_argument('--markdown',dest='md_path',type=Path); p.add_argument('--xlsx',dest='xlsx_path',type=Path); a=p.parse_args(argv)
 records,sources=load(a.records); report=build_inter_run_statistics(records,sources); rows=flatten_groups(report)
 if a.json_path: atomic(a.json_path,lambda x:x.write_text(json.dumps(report,indent=2,sort_keys=True,allow_nan=False)+'\n'))
 if a.csv_path: write_csv(a.csv_path,rows)
 if a.md_path: write_markdown(a.md_path,rows)
 if a.xlsx_path: write_xlsx(a.xlsx_path,report,rows)
 if not any((a.json_path,a.csv_path,a.md_path,a.xlsx_path)): print(json.dumps(report,indent=2))
 return 0
if __name__=='__main__': raise SystemExit(main())
