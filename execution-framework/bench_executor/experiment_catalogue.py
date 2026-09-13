"""Validate the frozen experiment catalogue and audit metric coverage."""
from __future__ import annotations

import csv
import json
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from bench_executor.rdf_experiment_manifest import system_adapter_specifications

CATALOGUE_SCHEMA = "krown-experiment-catalogue-v1"
AUDIT_SCHEMA = "krown-experiment-coverage-audit-v1"
VALID_STORAGE = {"file-backed", "in-memory", "persistent-database"}
VALID_METRICS = {
    "query_time", "query_peak_ram", "cold_load_or_parse_time",
    "warm_load_or_parse_time", "build_time", "build_peak_ram",
    "persistent_representation_size", "result_correctness",
    "failure_classification",
}
METRIC_STATUS = {"validated", "implemented-unvalidated", "missing", "not-applicable"}

# Conservative known coverage. A cell becomes validated only after a focused test.
KNOWN_COVERAGE = {
    "hdt-rdflib/optimized-in-memory": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"validated","warm_load_or_parse_time":"validated","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "pycottas/default": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"validated","warm_load_or_parse_time":"validated","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "vortex-rdf/dictionary-secondary-by-reference": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"validated","warm_load_or_parse_time":"validated","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "vortex-rdf/dictionary-secondary-by-reference-memory": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"validated","warm_load_or_parse_time":"validated","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "vortex-rdf/dictionary-secondary-by-copy": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"validated","warm_load_or_parse_time":"validated","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "vortex-rdf/dictionary-secondary-by-copy-memory": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"validated","warm_load_or_parse_time":"validated","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "comunica/hdt": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"validated","warm_load_or_parse_time":"validated","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "oxigraph/memory": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"not-applicable","warm_load_or_parse_time":"not-applicable","build_time":"not-applicable","build_peak_ram":"not-applicable","persistent_representation_size":"not-applicable","result_correctness":"validated","failure_classification":"validated"},
    "oxigraph/rocksdb": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"not-applicable","warm_load_or_parse_time":"not-applicable","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "rdflib/default": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"validated","warm_load_or_parse_time":"validated","build_time":"not-applicable","build_peak_ram":"not-applicable","persistent_representation_size":"not-applicable","result_correctness":"validated","failure_classification":"validated"},
    "fuseki/memory": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"not-applicable","warm_load_or_parse_time":"not-applicable","build_time":"not-applicable","build_peak_ram":"not-applicable","persistent_representation_size":"not-applicable","result_correctness":"validated","failure_classification":"validated"},
    "fuseki/tdb2": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"not-applicable","warm_load_or_parse_time":"not-applicable","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "qlever/default": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"not-applicable","warm_load_or_parse_time":"not-applicable","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
    "virtuoso/default": {"query_time":"validated","query_peak_ram":"validated","cold_load_or_parse_time":"not-applicable","warm_load_or_parse_time":"not-applicable","build_time":"validated","build_peak_ram":"validated","persistent_representation_size":"validated","result_correctness":"validated","failure_classification":"validated"},
}


def load_catalogue(path: Path) -> dict[str, Any]:
    value = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(value, Mapping) or value.get("schema") != CATALOGUE_SCHEMA:
        raise ValueError("unsupported experiment catalogue schema")
    required = {"schema", "catalogue_id", "scope", "required_metrics", "required_outputs", "oom_rule", "setups"}
    if set(value) != required:
        raise ValueError("catalogue has unexpected or missing fields")
    metrics = value["required_metrics"]
    if not isinstance(metrics, list) or set(metrics) != VALID_METRICS or len(metrics) != len(set(metrics)):
        raise ValueError("required_metrics must contain the complete unique metric contract")
    if value["required_outputs"] != ["json", "csv", "markdown", "xlsx"]:
        raise ValueError("required_outputs must freeze JSON, CSV, Markdown, and XLSX")
    if value["oom_rule"] != "confirmed-process-memory-exhaustion-only":
        raise ValueError("OOM rule is not strict")
    setups = value["setups"]
    if not isinstance(setups, list) or not setups:
        raise ValueError("setups must be a non-empty array")
    ids=set(); systems=set()
    for setup in setups:
        fields={"id","system_id","required","engine","representation","storage_mode","execution_interface","formats","notes"}
        if not isinstance(setup,Mapping) or set(setup)!=fields:
            raise ValueError("setup has unexpected or missing fields")
        if setup["id"] in ids: raise ValueError("duplicate setup ID")
        ids.add(setup["id"])
        if setup["system_id"] in systems: raise ValueError("duplicate system ID")
        systems.add(setup["system_id"])
        if setup["storage_mode"] not in VALID_STORAGE: raise ValueError("unsupported storage mode")
        if not isinstance(setup["required"],bool): raise TypeError("required must be boolean")
        if not isinstance(setup["formats"],list) or not setup["formats"]: raise ValueError("formats must be non-empty")
    return dict(value)


def build_coverage_audit(catalogue: Mapping[str, Any]) -> dict[str, Any]:
    registered={item.system_id for item in system_adapter_specifications()}
    rows=[]
    for setup in catalogue["setups"]:
        system=setup["system_id"]; coverage=KNOWN_COVERAGE.get(system,{})
        for metric in catalogue["required_metrics"]:
            status=coverage.get(metric,"missing")
            if status not in METRIC_STATUS: raise ValueError("invalid metric status")
            rows.append({"setup_id":setup["id"],"system_id":system,"required":setup["required"],"storage_mode":setup["storage_mode"],"metric":metric,"status":status})
    required_setups=[item for item in catalogue["setups"] if item["required"]]
    missing_systems=sorted(item["system_id"] for item in required_setups if item["system_id"] not in registered)
    blocking_cells=[row for row in rows if row["required"] and row["status"] in {"missing","implemented-unvalidated"}]
    return {"schema":AUDIT_SCHEMA,"catalogue_id":catalogue["catalogue_id"],"registered_system_ids":sorted(registered),"missing_required_system_ids":missing_systems,"rows":rows,"blocking_cells":blocking_cells,"ready_for_bsbm_10k":not missing_systems and not blocking_cells}


def write_audit_outputs(audit: Mapping[str,Any], output_root: Path) -> None:
    root=Path(output_root);root.mkdir(parents=True,exist_ok=True)
    (root/"coverage-audit.json").write_text(json.dumps(audit,indent=2,sort_keys=True)+"\n")
    with (root/"coverage-audit.csv").open("w",newline="",encoding="utf-8") as stream:
        writer=csv.DictWriter(stream,fieldnames=["setup_id","system_id","required","storage_mode","metric","status"]);writer.writeheader();writer.writerows(audit["rows"])
    lines=["# Experiment coverage audit","",f"Ready for BSBM 10k: **{str(audit['ready_for_bsbm_10k']).lower()}**","",f"Missing required systems: {len(audit['missing_required_system_ids'])}",f"Blocking metric cells: {len(audit['blocking_cells'])}","","## Missing systems",""]
    lines.extend([f"- `{item}`" for item in audit["missing_required_system_ids"]] or ["- None"])
    lines.extend(["","## Blocking cells",""])
    lines.extend([f"- `{row['setup_id']}`: `{row['metric']}` = `{row['status']}`" for row in audit["blocking_cells"]] or ["- None"])
    (root/"coverage-audit.md").write_text("\n".join(lines)+"\n")
    workbook = Workbook()
    coverage = workbook.active
    coverage.title = "System metric coverage"
    columns = ["setup_id", "system_id", "required", "storage_mode", "metric", "status"]
    coverage.append(columns)
    for row in audit["rows"]:
        coverage.append([row[name] for name in columns])
    readiness = workbook.create_sheet("Readiness")
    readiness.append(["field", "value"])
    readiness.append(["catalogue_id", audit["catalogue_id"]])
    readiness.append(["missing_required_system_count", len(audit["missing_required_system_ids"])])
    readiness.append(["blocking_metric_cell_count", len(audit["blocking_cells"])])
    readiness.append(["ready_for_bsbm_10k", audit["ready_for_bsbm_10k"]])
    fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in (coverage, readiness):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, name in enumerate([cell.value for cell in sheet[1]], 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(max(len(str(name)) + 2, 14), 42)
    temporary_xlsx = root / ".coverage-audit.xlsx.tmp"
    try:
        workbook.save(temporary_xlsx)
        temporary_xlsx.replace(root / "coverage-audit.xlsx")
    finally:
        temporary_xlsx.unlink(missing_ok=True)
