import json,tempfile,unittest
from pathlib import Path
from bench_executor.rdf_inter_run_statistics import build_inter_run_statistics,numeric_statistics
from analyze_rdf_inter_runs import main

def record(run,outcome='completed',elapsed=10,phase='measured',storage='file-backed',matrix=None):
 return {'benchmark':'BSBM','dataset':'10k','workload':'explore','experiment_id':'e','system':'s/default','representation':'rdf/source','query_id':'q1','phase':phase,'measurement_boundary':'query-and-full-materialization','execution_mode':{'storage':storage,'process_temperature':'warm-process','lifecycle':'shared'},'matrix_run_id':matrix or f'm{run}','run':run,'status':'ok' if outcome=='completed' else 'timeout','outcome':{'category':outcome},'elapsed_ns':elapsed}

class InterRunStatisticsTests(unittest.TestCase):
 def test_statistics_and_linear_percentiles(self):
  value=numeric_statistics([0,10,20]); self.assertEqual(value['mean'],10); self.assertEqual(value['sample_standard_deviation'],10); self.assertEqual(value['p90'],18)
 def test_non_completed_outcome_is_count_only(self):
  report=build_inter_run_statistics([record(0,elapsed=10),record(1,'timeout',60000000000)])
  group=report['groups'][0]; self.assertEqual(group['elapsed_ns']['observation_count'],1); self.assertEqual(group['elapsed_ns']['mean'],10); self.assertEqual(group['outcome_counts']['timeout'],1)
 def test_zero_differs_from_no_observation(self):
  self.assertEqual(numeric_statistics([0])['mean'],0); self.assertIsNone(numeric_statistics([])['mean'])
 def test_cold_warm_and_storage_do_not_merge(self):
  rows=[record(0),record(1,phase='warmup'),record(2,storage='in-memory')]
  self.assertEqual(build_inter_run_statistics(rows)['group_count'],3)
 def test_duplicate_identity_is_rejected(self):
  row=record(0)
  with self.assertRaisesRegex(ValueError,'duplicate'): build_inter_run_statistics([row,row])
 def test_all_report_formats(self):
  from openpyxl import load_workbook
  with tempfile.TemporaryDirectory() as directory:
   root=Path(directory); source=root/'records.jsonl'; source.write_text('\n'.join(json.dumps(record(i,elapsed=(i+1)*10)) for i in range(3))+'\n')
   outputs=[root/'r.json',root/'r.csv',root/'r.md',root/'r.xlsx']
   self.assertEqual(main([str(source),'--json',str(outputs[0]),'--csv',str(outputs[1]),'--markdown',str(outputs[2]),'--xlsx',str(outputs[3])]),0)
   self.assertTrue(all(p.stat().st_size for p in outputs)); self.assertEqual(load_workbook(outputs[3],read_only=True).sheetnames,['Inter-run statistics','Outcome counts','Observation provenance'])
if __name__=='__main__': unittest.main()
