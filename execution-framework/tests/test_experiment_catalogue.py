import json,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from bench_executor.experiment_catalogue import KNOWN_COVERAGE,build_coverage_audit,load_catalogue,write_audit_outputs

class Spec:
 def __init__(self,value):self.system_id=value
class CatalogueTests(unittest.TestCase):
 def catalogue(self):return load_catalogue(Path(__file__).resolve().parents[1]/"experiment-catalogue.json")
 def test_supervisor_scope_is_frozen(self):
  value=self.catalogue();self.assertEqual(value["scope"]["blocking_benchmarks"],["bsbm","krown-synthetic","dbbench"]);self.assertEqual(value["required_outputs"], ["json","csv","markdown","xlsx"]);self.assertEqual(value["oom_rule"],"confirmed-process-memory-exhaustion-only")
  systems={item["system_id"] for item in value["setups"]};self.assertIn("fuseki/memory",systems);self.assertIn("fuseki/tdb2",systems);self.assertIn("hdt-rdflib/optimized-in-memory",systems);self.assertIn("pyoxigraph-rdflib/memory",systems);self.assertIn("vortex-rdf/dictionary-secondary-by-reference-memory",systems)
 def test_missing_systems_and_cells_block_readiness(self):
  value=self.catalogue()
  with patch("bench_executor.experiment_catalogue.system_adapter_specifications",lambda:[Spec("pycottas/default")]):audit=build_coverage_audit(value)
  self.assertFalse(audit["ready_for_bsbm_10k"]);self.assertIn("fuseki/memory",audit["missing_required_system_ids"]);self.assertEqual(audit["blocking_cells"],[])
 def test_unresolved_metric_cell_blocks_readiness(self):
  with patch.dict(KNOWN_COVERAGE["qlever/default"],{"result_correctness":"implemented-unvalidated"}):
   audit=build_coverage_audit(self.catalogue())
  self.assertFalse(audit["ready_for_bsbm_10k"])
  self.assertEqual(audit["missing_required_system_ids"],[])
  blockers={(row["system_id"],row["metric"],row["status"]) for row in audit["blocking_cells"]}
  self.assertIn(("qlever/default","result_correctness","implemented-unvalidated"),blockers)
 def test_external_http_reopen_cells_are_not_applicable(self):
  audit=build_coverage_audit(self.catalogue())
  statuses={(row["system_id"],row["metric"]):row["status"] for row in audit["rows"]}
  systems=("fuseki/memory","fuseki/tdb2","qlever/default","oxigraph/memory","oxigraph/rocksdb","virtuoso/default")
  for system in systems:
   self.assertEqual(statuses[(system,"cold_load_or_parse_time")],"not-applicable")
   self.assertEqual(statuses[(system,"warm_load_or_parse_time")],"not-applicable")
  self.assertEqual(statuses[("qlever/default","result_correctness")],"validated")
  self.assertTrue(audit["ready_for_bsbm_10k"])
  self.assertEqual(audit["blocking_cells"],[])
  self.assertEqual(audit["missing_required_system_ids"],[])
 def test_outputs_are_written(self):
  value=self.catalogue()
  with patch("bench_executor.experiment_catalogue.system_adapter_specifications",lambda:[Spec(item["system_id"]) for item in value["setups"]]):audit=build_coverage_audit(value)
  with tempfile.TemporaryDirectory() as d:
   from openpyxl import load_workbook
   root=Path(d);write_audit_outputs(audit,root)
   for name in ("coverage-audit.json","coverage-audit.csv","coverage-audit.md","coverage-audit.xlsx"):self.assertTrue((root/name).is_file())
   self.assertEqual(load_workbook(root/"coverage-audit.xlsx",read_only=True).sheetnames,["System metric coverage","Readiness"])
if __name__=="__main__":unittest.main()
