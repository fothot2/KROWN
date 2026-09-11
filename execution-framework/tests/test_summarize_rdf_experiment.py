#!/usr/bin/env python3
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from summarize_rdf_experiment import build_report, main


def sample_summary():
    return {
        "status": "ok",
        "experiments": [{
            "system": "vortex/reference",
            "representation": "vortex/reference",
            "status": "ok", "record_count": 3, "success_count": 3, "failure_count": 0, "skipped_count": 0,
            "outcome_counts": {"completed": 3, "timeout": 0, "skipped": 0, "engine-error": 0, "semantic-mismatch": 0, "confirmed-oom": 0},
            "execution_mode": {"storage": "file-backed", "process_temperature": "warm-process", "lifecycle": "shared"},
            "resource_metrics": {"max_rss_kib": 2048, "user_cpu_ns": 3_000_000, "system_cpu_ns": 2_000_000, "major_page_faults": 1, "minor_page_faults": 2, "input_blocks": 3, "output_blocks": 4},
            "workload_timing": {"schema": "rdf-workload-timing-v1", "phases": {"warmup": {"attempt_count": 1, "attempt_total_ns": 4_000_000}, "measured": {"attempt_count": 2, "attempt_total_ns": 10_000_000, "successful_attempt_count": 2, "successful_attempt_total_ns": 10_000_000}}},
            "system_timing": {"schema": "rdf-system-timing-v1", "stages_ns": {"artifact_open_or_load": 2, "engine_startup": 0, "warmup": 4, "measured": 10, "engine_shutdown": 1, "validation": 2, "archive": 0, "preflight": 0, "unclassified": 1}, "stages_sum_ns": 20, "total_wall_ns": 20, "reconciled": True},
        }],
    }


class ReportTests(unittest.TestCase):
    def test_build_report_aggregates_timing_resources_and_modes(self):
        report = build_report(sample_summary())
        row = report["experiments"][0]
        self.assertEqual(row["measured_total_ms"], 10.0)
        self.assertEqual(row["measured_mean_ms"], 5.0)
        self.assertEqual(row["max_rss_mib"], 2.0)
        self.assertEqual(row["storage"], "file-backed")

    def test_unreconciled_system_timing_is_rejected(self):
        value = sample_summary()
        value["experiments"][0]["system_timing"]["total_wall_ns"] = 21
        with self.assertRaisesRegex(ValueError, "does not reconcile"):
            build_report(value)

    def test_partial_report_keeps_failures_out_of_successful_time(self):
        value = sample_summary(); experiment = value["experiments"][0]
        value["status"] = experiment["status"] = "completed_with_failures"
        experiment["failure_count"] = 1
        experiment["outcome_counts"].update({"completed": 1, "timeout": 1})
        measured = experiment["workload_timing"]["phases"]["measured"]
        measured.update({"successful_attempt_count": 1, "successful_attempt_total_ns": 3_000_000})
        row = build_report(value)["experiments"][0]
        self.assertEqual(row["successful_query_total_ms"], 3.0)
        self.assertEqual(row["timeout_count"], 1)

    def test_cli_writes_all_four_formats(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory); source=root/"summary.json"
            outputs=[root/"report.json",root/"report.csv",root/"report.md",root/"report.xlsx"]
            source.write_text(json.dumps(sample_summary()))
            self.assertEqual(main([str(source),"--json",str(outputs[0]),"--csv",str(outputs[1]),"--markdown",str(outputs[2]),"--xlsx",str(outputs[3])]),0)
            self.assertTrue(all(path.is_file() and path.stat().st_size for path in outputs))
            self.assertIn("confirmed_oom_count", outputs[2].read_text())
            self.assertIn("Experiment summary", load_workbook(outputs[3], read_only=True).sheetnames)

    def test_cli_writes_json_and_csv(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, json_path, csv_path = root/"summary.json", root/"report.json", root/"report.csv"
            source.write_text(json.dumps(sample_summary()))
            self.assertEqual(main([str(source), "--json", str(json_path), "--csv", str(csv_path)]), 0)
            self.assertEqual(json.loads(json_path.read_text())["experiment_count"], 1)
            self.assertIn("measured_total_ms", csv_path.read_text())


if __name__ == "__main__":
    unittest.main()
